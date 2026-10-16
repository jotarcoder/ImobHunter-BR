from __future__ import annotations

import csv
import html
import json
from pathlib import Path
from collections.abc import Mapping, Sequence
from typing import Any


MARKET_FIELDS = (
    "name",
    "cnpj",
    "cnae",
    "address",
    "city",
    "state",
    "latitude",
    "longitude",
    "distance_km",
    "revenue_reported",
    "revenue_period",
    "revenue_estimate",
    "estimate_method",
    "revenue_status",
    "source",
    "updated_at",
)


def export_market_results(
    results: Sequence[Mapping[str, Any]],
    path: str | Path,
    query: Mapping[str, Any] | None = None,
) -> Path:
    """Save a market-intelligence search in CSV, XLSX, JSON, or standalone HTML."""
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix not in {".csv", ".xlsx", ".json", ".html", ".htm"}:
        raise ValueError("Formato inválido. Use CSV, XLSX, JSON ou HTML.")

    rows = [
        {field: item.get(field, "") for field in MARKET_FIELDS}
        for item in results
    ]
    query_data = dict(query or {})

    if suffix == ".csv":
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=MARKET_FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
    elif suffix == ".json":
        payload = {
            "title": "ImobHunter — Inteligência de Mercado",
            "query": query_data,
            "count": len(rows),
            "records": rows,
        }
        target.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False),
            encoding="utf-8",
        )
    elif suffix in {".html", ".htm"}:
        headers = "".join(f"<th>{html.escape(field)}</th>" for field in MARKET_FIELDS)
        body = "".join(
            "<tr>" + "".join(
                f"<td>{html.escape(str(row.get(field, '') if row.get(field) is not None else ''))}</td>"
                for field in MARKET_FIELDS
            ) + "</tr>"
            for row in rows
        )
        query_summary = " · ".join(
            f"{html.escape(str(key))}: {html.escape(str(value))}"
            for key, value in query_data.items()
        )
        report = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>ImobHunter — Inteligência de Mercado</title>
<style>
body{{font:14px system-ui,sans-serif;background:#07111b;color:#eaf2f7;margin:28px}}
h1{{color:#35ff77}}.meta{{color:#a9bdcb;margin-bottom:18px}}
.table-wrap{{overflow:auto;border:1px solid #254156;border-radius:10px}}
table{{border-collapse:collapse;min-width:100%;background:#0b1926}}
th,td{{padding:9px 11px;text-align:left;border-bottom:1px solid #254156;white-space:nowrap}}
th{{background:#102332;color:#35ff77;position:sticky;top:0}}
tr:nth-child(even){{background:#0d1d2b}}
</style></head><body>
<h1>Inteligência de Mercado</h1>
<div class="meta">Registros: {len(rows)} · {query_summary}</div>
<div class="table-wrap"><table><thead><tr>{headers}</tr></thead><tbody>{body}</tbody></table></div>
<p class="meta">ImobHunter-BR · Os valores ausentes não foram estimados automaticamente.</p>
</body></html>"""
        target.write_text(report, encoding="utf-8")
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inteligência de Mercado"
        sheet.append(list(MARKET_FIELDS))
        for row in rows:
            sheet.append([row.get(field, "") for field in MARKET_FIELDS])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="35FF77")
            cell.fill = PatternFill("solid", fgColor="102332")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, field in enumerate(MARKET_FIELDS, 1):
            max_length = max(
                len(field),
                *(
                    len(str(sheet.cell(row, index).value or ""))
                    for row in range(2, sheet.max_row + 1)
                ),
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max_length + 2, 42)
        workbook.save(target)

    return target
